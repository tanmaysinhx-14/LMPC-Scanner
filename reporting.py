from __future__ import annotations

import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO
import json
from pathlib import Path
import textwrap
from typing import Any
from xml.sax.saxutils import escape

import rule_config


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _jsonable(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (AttributeError, TypeError, ValueError):
            return str(value)
    return value


def _results_list(scan: dict[str, Any]) -> list[dict[str, Any]]:
    results = scan.get("results", [])
    if isinstance(results, dict):
        return [
            {
                "rule_class": rule_class,
                **details,
            }
            for rule_class, details in results.items()
            if isinstance(details, dict)
        ]
    return [dict(result) for result in results if isinstance(result, dict)]


def _image_paths(scan: dict[str, Any]) -> list[str]:
    paths = scan.get("image_paths")
    if paths is None:
        paths = scan.get("image_path", [])
    if isinstance(paths, (str, Path)):
        paths = [paths]
    return [str(path) for path in paths or [] if str(path).strip()]


def _rule_config_block(scan: dict[str, Any], results: list[dict[str, Any]]) -> dict[str, Any]:
    """Describe the rule set this report was produced under.

    Without this a reader cannot tell a clean pack from a narrow profile, so it
    is part of the report rather than an admin-only detail.
    """

    config = rule_config.normalize_config(scan.get("rule_config"))
    assessed = {
        rule_config.result_metadata(result)["rule_class"]
        for result in results
        if rule_config.result_metadata(result)["in_scope"]
    }
    in_scope: list[dict[str, Any]] = []
    out_of_scope: list[dict[str, Any]] = []
    for entry in rule_config.RULE_CATALOGUE:
        rule_class = entry["rule_class"]
        setting = config["rules"][rule_class]
        row = {
            "rule_class": rule_class,
            "declaration": entry["label"],
            "statute": setting["statute"],
            "severity": setting["severity"],
            "penalty_if_failed": int(setting["penalty"]),
        }
        (in_scope if setting["enabled"] else out_of_scope).append(row)
    return {
        "profile_name": config["profile_name"],
        "config_version": rule_config.config_version(config),
        "schema_version": config["schema_version"],
        "rules_in_scope": in_scope,
        "rules_out_of_scope": out_of_scope,
        "assessed_rule_classes": sorted(assessed),
    }


def build_report_payload(
    scan: dict[str, Any],
    generated_by: str = "",
) -> dict[str, Any]:
    results = _results_list(scan)
    for result in results:
        metadata = rule_config.result_metadata(result)
        result.setdefault("severity", metadata["severity"])
        result.setdefault("rule_label", metadata["label"])
        result.setdefault("statute", metadata["statute"])
        result.setdefault("in_scope", metadata["in_scope"])
    summary = rule_config.summarize(results)
    paths = _image_paths(scan)
    hashes = scan.get("evidence_hashes", [])
    if isinstance(hashes, str):
        try:
            hashes = json.loads(hashes)
        except json.JSONDecodeError:
            hashes = [hashes]
    return {
        "report_type": "LMPC Rule 6 inspection report",
        "generated_at": _now(),
        "generated_by": generated_by,
        "scan": {
            "scan_id": scan.get("scan_id"),
            "status": scan.get("status", "READY"),
            "timestamp": scan.get("timestamp", ""),
            "inspector_id": scan.get("inspector_id"),
            "inspector_username": scan.get("inspector_username", ""),
            "product_name": scan.get("product_name", ""),
            "evidence_notes": scan.get("evidence_notes", ""),
            "reviewer_id": scan.get("reviewer_id"),
            "reviewed_at": scan.get("reviewed_at", ""),
            "review_reason": scan.get("review_reason", ""),
        },
        "rule_config": _rule_config_block(scan, results),
        "summary": {
            "total_fields": summary["total"],
            "assessed_fields": summary["assessed"],
            "skipped_fields": summary["skipped"],
            "passed_fields": summary["passed"],
            "failed_fields": summary["failed"],
            "estimated_penalty": summary["penalty"],
            "highest_severity": summary["highest_severity"],
            "failed_by_severity": summary["failed_by_severity"],
        },
        "evidence": {
            "image_paths": paths,
            "image_hashes": list(hashes) if isinstance(hashes, list) else [],
        },
        "results": _jsonable(results),
    }


def generate_json_report(scan: dict[str, Any], generated_by: str = "") -> bytes:
    payload = build_report_payload(scan, generated_by)
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def generate_csv_report(scan: dict[str, Any], generated_by: str = "") -> bytes:
    payload = build_report_payload(scan, generated_by)
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(
        [
            "scan_id",
            "status",
            "timestamp",
            "inspector",
            "rule_profile",
            "config_version",
            "rule_class",
            "declaration",
            "in_scope",
            "severity",
            "statute",
            "is_compliant",
            "extracted_text",
            "reason",
            "parsed_data",
            "penalty_amount",
            "evidence_images",
        ]
    )
    scan_data = payload["scan"]
    profile = payload["rule_config"]
    evidence_images = " | ".join(payload["evidence"]["image_paths"])
    for result in payload["results"]:
        metadata = rule_config.result_metadata(result)
        writer.writerow(
            [
                scan_data.get("scan_id", ""),
                scan_data.get("status", ""),
                scan_data.get("timestamp", ""),
                scan_data.get("inspector_username", ""),
                profile["profile_name"],
                profile["config_version"],
                result.get("rule_class", ""),
                metadata["label"],
                "YES" if metadata["in_scope"] else "NO",
                metadata["severity"],
                metadata["statute"],
                ("PASS" if result.get("is_compliant", False) else "FAIL")
                if metadata["in_scope"]
                else "NOT ASSESSED",
                result.get("extracted_text", ""),
                result.get("reason", ""),
                json.dumps(
                    {
                        key: value
                        for key, value in result.items()
                        if key
                        not in {"rule_class", "extracted_text", "is_compliant", "penalty_amount", "reason"}
                    },
                    ensure_ascii=False,
                    default=str,
                ),
                int(result.get("penalty_amount", 0) or 0),
                evidence_images,
            ]
        )
    return output.getvalue().encode("utf-8-sig")


def generate_docx_report(scan: dict[str, Any], generated_by: str = "") -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.shared import Inches, Pt
    except ImportError as error:
        raise RuntimeError("Editable DOCX reports require python-docx") from error

    payload = build_report_payload(scan, generated_by)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)

    document.add_heading("LMPC Rule 6 Inspection Report", level=0)
    document.add_paragraph(
        f"Generated: {payload['generated_at']} | Generated by: {generated_by or 'Scanner'}"
    )
    scan_data = payload["scan"]
    summary = payload["summary"]
    profile = payload["rule_config"]
    metadata = document.add_table(rows=0, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.style = "Light Shading Accent 1"
    for label, value in (
        ("Scan ID", scan_data.get("scan_id", "")),
        ("Status", scan_data.get("status", "")),
        ("Inspector", scan_data.get("inspector_username", "")),
        ("Product", scan_data.get("product_name", "") or "Not detected"),
        ("Captured", scan_data.get("timestamp", "")),
        ("Rule profile", f"{profile['profile_name']} ({profile['config_version']})"),
        ("Declarations assessed", summary["assessed_fields"]),
        ("Not assessed (out of scope)", summary["skipped_fields"]),
        ("Fields passed", summary["passed_fields"]),
        ("Fields failed", summary["failed_fields"]),
        ("Highest severity", summary["highest_severity"] or "None"),
        ("Estimated penalty", f"INR {summary['estimated_penalty']:,}"),
    ):
        cells = metadata.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)

    document.add_heading("Compliance results", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, heading in zip(
        table.rows[0].cells,
        ("Declaration", "Statutory basis", "Severity", "Status", "Reason", "Penalty"),
    ):
        cell.text = heading
    for result in payload["results"]:
        row_meta = rule_config.result_metadata(result)
        cells = table.add_row().cells
        cells[0].text = row_meta["label"]
        cells[1].text = row_meta["statute"]
        cells[2].text = row_meta["severity"]
        cells[3].text = (
            ("PASS" if result.get("is_compliant", False) else "FAIL")
            if row_meta["in_scope"]
            else "NOT ASSESSED"
        )
        cells[4].text = (
            f"{str(result.get('reason', ''))}\n"
            f"Read from pack: {str(result.get('extracted_text', '') or 'No text detected')}"
        )
        cells[5].text = f"INR {int(result.get('penalty_amount', 0) or 0):,}"

    document.add_heading("Rule set applied", level=1)
    document.add_paragraph(
        f"Profile '{profile['profile_name']}', config fingerprint {profile['config_version']}. "
        "Re-running this scan against the same fingerprint reproduces this report."
    )
    if profile["rules_out_of_scope"]:
        document.add_paragraph(
            "Declarations excluded by this profile and therefore not assessed: "
            + ", ".join(row["declaration"] for row in profile["rules_out_of_scope"])
        )
    else:
        document.add_paragraph("No declaration was excluded; the full Rule 6 set was assessed.")

    notes = str(scan_data.get("evidence_notes", "") or "").strip()
    review_reason = str(scan_data.get("review_reason", "") or "").strip()
    if notes or review_reason:
        document.add_heading("Workflow notes", level=1)
        if notes:
            document.add_paragraph(f"Inspector evidence notes: {notes}")
        if review_reason:
            document.add_paragraph(f"Verifier decision note: {review_reason}")

    document.add_heading("Attached evidence", level=1)
    image_paths = payload["evidence"]["image_paths"]
    hashes = payload["evidence"]["image_hashes"]
    for index, path in enumerate(image_paths):
        document.add_paragraph(f"Evidence {index + 1}: {path}")
        if index < len(hashes) and hashes[index]:
            document.add_paragraph(f"SHA-256: {hashes[index]}")
        if Path(path).is_file():
            try:
                document.add_picture(path, width=Inches(5.8))
            except (OSError, ValueError):
                document.add_paragraph("The image could not be embedded; see the stored path above.")
        else:
            document.add_paragraph("The stored image is unavailable at report generation time.")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _wrap_lines(value: Any, width: int = 92) -> list[str]:
    text = " ".join(str(value or "").split())
    if not text:
        return [""]
    return textwrap.wrap(text, width=width, break_long_words=False, break_on_hyphens=False) or [""]


def _reportlab_text(value: Any) -> str:
    return escape(str(value or ""))


def _generate_matplotlib_pdf(payload: dict[str, Any]) -> bytes:
    try:
        import matplotlib

        matplotlib.use("Agg")
        from matplotlib.backends.backend_pdf import PdfPages
        from matplotlib.figure import Figure
        from PIL import Image
    except ImportError as error:
        raise RuntimeError("PDF reports require matplotlib or reportlab") from error

    output = BytesIO()
    with PdfPages(output) as pdf:
        profile = payload["rule_config"]
        lines: list[str] = [
            "LMPC Rule 6 Inspection Report",
            f"Generated: {payload['generated_at']}",
            f"Scan ID: {payload['scan'].get('scan_id', 'Draft')}",
            f"Status: {payload['scan'].get('status', 'READY')}",
            f"Inspector: {payload['scan'].get('inspector_username', '') or 'Unknown'}",
            f"Product: {payload['scan'].get('product_name', '') or 'Not detected'}",
            f"Rule profile: {profile['profile_name']} (config {profile['config_version']})",
            "",
            f"Declarations: {payload['summary']['assessed_fields']} assessed, "
            f"{payload['summary']['skipped_fields']} out of scope",
            f"Assessed: {payload['summary']['passed_fields']} passed, "
            f"{payload['summary']['failed_fields']} failed",
            f"Highest severity: {payload['summary']['highest_severity'] or 'None'}",
            f"Estimated penalty: INR {payload['summary']['estimated_penalty']:,}",
            "",
        ]
        for result in payload["results"]:
            row_meta = rule_config.result_metadata(result)
            status = (
                ("PASS" if result.get("is_compliant", False) else "FAIL")
                if row_meta["in_scope"]
                else "NOT ASSESSED"
            )
            lines.extend(
                [
                    f"{row_meta['label']}: {status} [{row_meta['severity']}]",
                    f"Statutory basis: {row_meta['statute']}",
                    f"Extracted: {result.get('extracted_text', '') or 'No text detected'}",
                    f"Reason: {result.get('reason', '')}",
                    f"Penalty: INR {int(result.get('penalty_amount', 0) or 0):,}",
                    "",
                ]
            )
        if profile["rules_out_of_scope"]:
            lines.extend([
                "Excluded by this profile (not assessed):",
                ", ".join(row["declaration"] for row in profile["rules_out_of_scope"]),
                "",
            ])
        if payload["scan"].get("evidence_notes"):
            lines.extend(["Inspector evidence notes:", str(payload["scan"]["evidence_notes"]), ""])
        if payload["scan"].get("review_reason"):
            lines.extend(["Verifier decision note:", str(payload["scan"]["review_reason"]), ""])
        lines.extend(["Attached evidence:", *payload["evidence"]["image_paths"]])

        page_lines: list[str] = []
        for line in lines:
            page_lines.extend(_wrap_lines(line))
        for offset in range(0, len(page_lines), 46):
            fig = Figure(figsize=(8.27, 11.69))
            axis = fig.add_axes([0.07, 0.06, 0.86, 0.88])
            axis.axis("off")
            axis.text(
                0,
                1,
                "\n".join(page_lines[offset : offset + 46]),
                va="top",
                ha="left",
                fontsize=9,
                family="DejaVu Sans",
            )
            pdf.savefig(fig)

        for index, path in enumerate(payload["evidence"]["image_paths"]):
            image_path = Path(path)
            if not image_path.is_file():
                continue
            try:
                image = Image.open(image_path)
                fig = Figure(figsize=(8.27, 11.69))
                axis = fig.add_axes([0.05, 0.08, 0.9, 0.84])
                axis.imshow(image)
                axis.axis("off")
                axis.set_title(f"Evidence image {index + 1}: {image_path.name}")
                pdf.savefig(fig)
                image.close()
            except (OSError, ValueError):
                continue
    return output.getvalue()


def generate_pdf_report(scan: dict[str, Any], generated_by: str = "") -> bytes:
    payload = build_report_payload(scan, generated_by)
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        return _generate_matplotlib_pdf(payload)

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph("LMPC Rule 6 Inspection Report", styles["Title"])]
    story.append(Paragraph(f"Generated: {_reportlab_text(payload['generated_at'])}", styles["Normal"]))
    story.append(Spacer(1, 10))
    scan_data = payload["scan"]
    summary = payload["summary"]
    profile = payload["rule_config"]
    metadata = [
        ["Scan ID", str(scan_data.get("scan_id", "Draft"))],
        ["Status", str(scan_data.get("status", "READY"))],
        ["Inspector", str(scan_data.get("inspector_username", "") or "Unknown")],
        ["Product", str(scan_data.get("product_name", "") or "Not detected")],
        ["Rule profile", f"{profile['profile_name']} (config {profile['config_version']})"],
        ["Assessed / out of scope", f"{summary['assessed_fields']} / {summary['skipped_fields']}"],
        ["Passed / failed", f"{summary['passed_fields']} / {summary['failed_fields']}"],
        ["Highest severity", str(summary["highest_severity"] or "None")],
        ["Estimated penalty", f"INR {summary['estimated_penalty']:,}"],
    ]
    metadata_table = Table(metadata, colWidths=[1.5 * inch, 5.2 * inch])
    metadata_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e8eef7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.extend([metadata_table, Spacer(1, 14), Paragraph("Compliance results", styles["Heading2"])])
    cell = ParagraphStyle("lmpc_cell", parent=styles["BodyText"], fontSize=7, leading=8.5, spaceBefore=0, spaceAfter=0)
    rows = [["Declaration", "Statutory basis", "Sev.", "Status", "Reason and extracted text", "Penalty"]]
    for result in payload["results"]:
        row_meta = rule_config.result_metadata(result)
        status = (
            ("PASS" if result.get("is_compliant", False) else "FAIL")
            if row_meta["in_scope"]
            else "N/A"
        )
        rows.append([
            Paragraph(_reportlab_text(row_meta["label"]), cell),
            Paragraph(_reportlab_text(row_meta["statute"]), cell),
            Paragraph(_reportlab_text(row_meta["severity"]), cell),
            status,
            Paragraph(
                f"{_reportlab_text(result.get('reason', ''))}<br/><b>Read:</b> "
                f"{_reportlab_text(result.get('extracted_text', '') or 'No text detected')}",
                cell,
            ),
            f"INR {int(result.get('penalty_amount', 0) or 0):,}",
        ])
    result_table = Table(
        rows,
        colWidths=[1.15 * inch, 1.45 * inch, 0.45 * inch, 0.55 * inch, 2.5 * inch, 0.6 * inch],
        repeatRows=1,
    )
    result_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(result_table)
    story.extend([Spacer(1, 10), Paragraph("Rule set applied", styles["Heading2"])])
    story.append(Paragraph(
        f"Profile '{_reportlab_text(profile['profile_name'])}', config fingerprint "
        f"{_reportlab_text(profile['config_version'])}. Re-running this scan against the same "
        "fingerprint reproduces this report.",
        styles["Normal"],
    ))
    excluded = ", ".join(row["declaration"] for row in profile["rules_out_of_scope"])
    story.append(Paragraph(
        f"Excluded by this profile and not assessed: {_reportlab_text(excluded)}."
        if excluded
        else "No declaration was excluded; the full Rule 6 set was assessed.",
        styles["Normal"],
    ))
    if scan_data.get("evidence_notes"):
        story.extend([Spacer(1, 10), Paragraph("Inspector evidence notes", styles["Heading2"]), Paragraph(_reportlab_text(scan_data["evidence_notes"]), styles["Normal"])])
    if scan_data.get("review_reason"):
        story.extend([Spacer(1, 10), Paragraph("Verifier decision note", styles["Heading2"]), Paragraph(_reportlab_text(scan_data["review_reason"]), styles["Normal"])])
    story.extend([Spacer(1, 10), Paragraph("Attached evidence", styles["Heading2"])])
    for index, path in enumerate(payload["evidence"]["image_paths"]):
        story.append(Paragraph(f"Evidence {index + 1}: {_reportlab_text(path)}", styles["Normal"]))
        if Path(path).is_file():
            try:
                story.append(Image(path, width=6.2 * inch, height=4.2 * inch, kind="proportional"))
                story.append(Spacer(1, 8))
            except (OSError, ValueError):
                continue
    document.build(story)
    return output.getvalue()


def report_filename(scan: dict[str, Any], extension: str) -> str:
    scan_id = scan.get("scan_id", "draft")
    return f"lmpc_scan_{scan_id}.{extension.lstrip('.') }"
